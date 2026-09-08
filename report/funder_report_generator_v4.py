import json
from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PROJECT PATHS
# ============================================================

ROOT = Path(__file__).resolve().parent.parent

KPI_DIR = ROOT / "output" / "funder_kpis"
INSIGHT_DIR = ROOT / "output" / "funder_insights"

TRANSLATED_FILE = ROOT / "output" / "translated_report.json"

OUTPUT_DIR = ROOT / "output" / "funder_reports_v4"

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


# ============================================================
# REPORT CONFIGURATION
# ============================================================

REPORT_PERIOD = "January 2025 – December 2025"

FUNDER_NAMES = {
    "ACS": "ACS",
    "AVAYA": "Avaya",
    "BBD": "BBD",
    "BREMBO": "Brembo Brake India Pvt. Ltd.",
    "DSSF": "DSSF",
    "GHATKOPAR": "Ghatkopar",
    "IDRF": "IDRF",
    "NICE": "NICE",
    "YARDI": "Yardi Software",
}


# ============================================================
# LOAD JSON
# ============================================================

def load_json(path):

    if not path.exists():
        raise FileNotFoundError(
            f"Required file not found:\n{path}"
        )

    with open(
        path,
        "r",
        encoding="utf-8"
    ) as f:

        return json.load(f)


def load_funder_data(funder):

    kpi_file = KPI_DIR / f"{funder}_kpis.json"
    insight_file = INSIGHT_DIR / f"{funder}_insights.json"

    kpis = load_json(kpi_file)
    insights = load_json(insight_file)

    return kpis, insights


def load_translated_report():

    return load_json(
        TRANSLATED_FILE
    )


# ============================================================
# SAFE JSON HELPERS
# ============================================================

def find_value(data, possible_keys, default=None):

    if not isinstance(data, dict):
        return default

    for key in possible_keys:

        if key in data:
            return data[key]

    return default


def recursive_find(data, target_keys):

    """
    Search nested dictionaries/lists for a matching key.
    Used because existing KPI JSON schemas may contain
    slightly different nesting.
    """

    if isinstance(data, dict):

        for key, value in data.items():

            normalized = str(key).lower().replace(
                " ",
                "_"
            )

            if normalized in target_keys:

                return value

            result = recursive_find(
                value,
                target_keys
            )

            if result is not None:
                return result

    elif isinstance(data, list):

        for item in data:

            result = recursive_find(
                item,
                target_keys
            )

            if result is not None:
                return result

    return None


def get_number(
    data,
    keys,
    default=0
):

    value = recursive_find(
        data,
        {
            str(k).lower().replace(" ", "_")
            for k in keys
        }
    )

    if value is None:
        return default

    try:
        return int(float(value))
    except:
        return value


def get_percent(
    data,
    keys,
    default=0
):

    value = recursive_find(
        data,
        {
            str(k).lower().replace(" ", "_")
            for k in keys
        }
    )

    if value is None:
        return default

    try:
        value = float(value)

        if value <= 1:
            value *= 100

        return round(
            value,
            2
        )

    except:
        return value


# ============================================================
# DOCUMENT FORMATTING
# ============================================================

def set_cell_shading(
    cell,
    fill
):

    tcPr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement(
        "w:shd"
    )

    shd.set(
        qn("w:fill"),
        fill
    )

    tcPr.append(
        shd
    )


def set_cell_text(
    cell,
    text,
    bold=False
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        str(text)
    )

    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


def add_heading(
    document,
    text,
    level=1
):

    paragraph = document.add_paragraph()

    paragraph.style = (
        f"Heading {level}"
    )

    run = paragraph.add_run(
        text
    )

    return paragraph


def add_body(
    document,
    text,
    bold_start=None
):

    if not text:
        return

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(6)

    if bold_start and text.startswith(
        bold_start
    ):

        run = paragraph.add_run(
            bold_start
        )

        run.bold = True

        paragraph.add_run(
            text[len(bold_start):]
        )

    else:

        paragraph.add_run(
            text
        )


def add_bullets(
    document,
    items
):

    if not items:
        return

    for item in items:

        if not item:
            continue

        paragraph = document.add_paragraph(
            style="List Bullet"
        )

        paragraph.add_run(
            str(item)
        )


def add_kpi_table(
    document,
    rows
):

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    set_cell_text(
        table.rows[0].cells[0],
        "Indicator",
        True
    )

    set_cell_text(
        table.rows[0].cells[1],
        "Value",
        True
    )

    for cell in table.rows[0].cells:
        set_cell_shading(
            cell,
            "D9EAF7"
        )

    for label, value in rows:

        cells = table.add_row().cells

        set_cell_text(
            cells[0],
            label
        )

        set_cell_text(
            cells[1],
            value
        )

    document.add_paragraph()

    return table


def add_activity_table(
    document,
    rows
):

    table = document.add_table(
        rows=1,
        cols=5
    )

    table.style = "Table Grid"

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    headers = [
        "Activity",
        "Total",
        "Continued",
        "Dropout",
        "Migrated"
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            True
        )

        set_cell_shading(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    for row in rows:

        cells = table.add_row().cells

        for i, value in enumerate(row):

            set_cell_text(
                cells[i],
                value
            )

    document.add_paragraph()

    return table


# ============================================================
# STORY HELPERS
# ============================================================

def get_story(
    translated,
    activity
):

    activities = translated.get(
        "activity_stories",
        {}
    )

    return activities.get(
        activity,
        {}
    )


def story_paragraph(
    document,
    activity_story
):

    if not activity_story:
        return

    description = activity_story.get(
        "description",
        ""
    )

    changes = activity_story.get(
        "observed_changes",
        []
    )

    context = activity_story.get(
        "implementation_context",
        []
    )

    challenges = activity_story.get(
        "challenges",
        []
    )

    if description:

        add_body(
            document,
            description
        )

    if changes:

        add_body(
            document,
            "Observed changes:"
        )

        add_bullets(
            document,
            changes
        )

    if context:

        add_body(
            document,
            "Implementation context:"
        )

        add_bullets(
            document,
            context
        )

    if challenges:

        add_body(
            document,
            "Implementation considerations:"
        )

        add_bullets(
            document,
            challenges
        )


# ============================================================
# COVER PAGE
# ============================================================

def create_cover(
    document,
    funder
):

    section = document.sections[0]

    section.top_margin = Inches(
        0.7
    )

    section.bottom_margin = Inches(
        0.7
    )

    section.left_margin = Inches(
        0.8
    )

    section.right_margin = Inches(
        0.8
    )

    document.add_paragraph(
        "\n"
    )

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        "DOOR STEP SCHOOL FOUNDATION"
    )

    r.bold = True
    r.font.size = Pt(18)

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        "Community Education Programme"
    )

    r.bold = True
    r.font.size = Pt(22)

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        f"supported by {FUNDER_NAMES.get(funder, funder)}"
    )

    r.bold = True
    r.font.size = Pt(17)

    document.add_paragraph(
        "\n"
    )

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        "PROJECT PROGRESS REPORT"
    )

    r.bold = True
    r.font.size = Pt(26)

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        REPORT_PERIOD
    )

    r.bold = True
    r.font.size = Pt(18)

    document.add_paragraph(
        "\n\n\n"
    )

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run(
        "Community Education Programme\n"
        "Foundational Literacy • Numeracy • "
        "Learning Continuity • Community Engagement"
    )

    r.font.size = Pt(13)

    document.add_page_break()


# ============================================================
# INTRODUCTION
# ============================================================

def add_introduction(
    document,
    translated
):

    add_heading(
        document,
        "Introduction",
        1
    )

    overview = translated.get(
        "programme_overview",
        {}
    )

    objective = overview.get(
        "objective",
        ""
    )

    context = overview.get(
        "key_context",
        ""
    )

    add_body(
        document,
        "Door Step School Foundation works to address "
        "educational gaps among children from marginalised "
        "and migrant communities through innovative and "
        "engaging education programmes."
    )

    add_body(
        document,
        "The Community Education Programme focuses on "
        "foundational literacy and numeracy, learning "
        "continuity, school readiness and the creation "
        "of a supportive learning environment."
    )

    if objective:

        add_body(
            document,
            objective
        )

    if context:

        add_body(
            document,
            context
        )


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def add_executive_summary(
    document,
    funder,
    kpis,
    translated
):

    add_heading(
        document,
        "Executive Summary",
        1
    )

    overview = translated.get(
        "programme_overview",
        {}
    )

    target = overview.get(
        "target_community",
        ""
    )

    context = overview.get(
        "key_context",
        ""
    )

    if target:

        add_body(
            document,
            target
        )

    if context:

        add_body(
            document,
            context
        )

    total = get_number(
        kpis,
        [
            "total",
            "total_children",
            "children",
            "total_children_reached"
        ]
    )

    continued = get_number(
        kpis,
        [
            "continued",
            "continuation_count"
        ]
    )

    dropout = get_number(
        kpis,
        [
            "dropout",
            "dropout_count"
        ]
    )

    migrated = get_number(
        kpis,
        [
            "migrated",
            "migration_count"
        ]
    )

    continuation_rate = get_percent(
        kpis,
        [
            "continuation_rate",
            "continuation"
        ]
    )

    dropout_rate = get_percent(
        kpis,
        [
            "dropout_rate"
        ]
    )

    migration_rate = get_percent(
        kpis,
        [
            "migration_rate"
        ]
    )

    add_body(
        document,
        (
            f"Through the Community Education Programme "
            f"supported by {FUNDER_NAMES.get(funder, funder)}, "
            f"{total:,} children were reached during the "
            f"reporting period. The programme worked to "
            f"strengthen foundational learning while "
            f"supporting children's continued engagement "
            f"with education."
        )
    )

    add_kpi_table(
        document,
        [
            (
                "Children reached",
                f"{total:,}"
            ),
            (
                "Children continuing",
                f"{continued:,}"
            ),
            (
                "Continuation rate",
                f"{continuation_rate:.2f}%"
            ),
            (
                "Dropout",
                f"{dropout:,}"
            ),
            (
                "Dropout rate",
                f"{dropout_rate:.2f}%"
            ),
            (
                "Migrated",
                f"{migrated:,}"
            ),
            (
                "Migration rate",
                f"{migration_rate:.2f}%"
            ),
        ]
    )


# ============================================================
# HIGHLIGHTS
# ============================================================

def add_highlights(
    document,
    kpis,
    translated
):

    add_heading(
        document,
        "Highlights and Reach",
        1
    )

    total = get_number(
        kpis,
        [
            "total",
            "total_children",
            "children"
        ]
    )

    continuation = get_percent(
        kpis,
        [
            "continuation_rate",
            "continuation"
        ]
    )

    attendance_high = get_percent(
        kpis,
        [
            "high_attendance_rate",
            "attendance_high_rate"
        ]
    )

    add_bullets(
        document,
        [
            f"{total:,} children reached through the programme.",
            f"{continuation:.2f}% overall continuation during the reporting period.",
            (
                f"{attendance_high:.2f}% of children were in the "
                "high-attendance category."
                if attendance_high
                else ""
            ),
        ]
    )

    language = translated.get(
        "learning_outcomes",
        {}
    ).get(
        "language",
        {}
    )

    mathematics = translated.get(
        "learning_outcomes",
        {}
    ).get(
        "mathematics",
        {}
    )

    if language.get(
        "observed_progress"
    ):

        add_body(
            document,
            "Language learning context:"
        )

        add_bullets(
            document,
            language[
                "observed_progress"
            ]
        )

    if mathematics.get(
        "observed_progress"
    ):

        add_body(
            document,
            "Mathematics learning context:"
        )

        add_bullets(
            document,
            mathematics[
                "observed_progress"
            ]
        )


# ============================================================
# ACTIVITY-WISE SECTION
# ============================================================

def add_activity_section(
    document,
    kpis,
    translated
):

    add_heading(
        document,
        "Activity-wise Objectives, Reach and Observations",
        1
    )

    activity_rows = []

    activities = [
        (
            "Library Class",
            "library_class",
            "library_class"
        ),
        (
            "Study Class",
            "study_class",
            "study_class"
        ),
        (
            "Balwadi",
            "balwadi",
            "balwadi"
        ),
        (
            "Home Lending",
            "home_lending",
            "home_lending"
        ),
    ]

    for display_name, kpi_name, story_name in activities:

        # Find activity data recursively
        activity_data = recursive_find(
            kpis,
            {
                kpi_name,
                kpi_name.replace(
                    "_",
                    " "
                )
            }
        )

        if not isinstance(
            activity_data,
            dict
        ):

            activity_data = {}

        total = get_number(
            activity_data,
            [
                "total",
                "children",
                "count"
            ]
        )

        continued = get_number(
            activity_data,
            [
                "continued",
                "continuation_count"
            ]
        )

        dropout = get_number(
            activity_data,
            [
                "dropout",
                "dropout_count"
            ]
        )

        migrated = get_number(
            activity_data,
            [
                "migrated",
                "migration_count"
            ]
        )

        if total:

            activity_rows.append(
                [
                    display_name,
                    total,
                    continued,
                    dropout,
                    migrated
                ]
            )

    if activity_rows:

        add_activity_table(
            document,
            activity_rows
        )

    # --------------------------------------------------------
    # Narrative activity sections
    # --------------------------------------------------------

    for display_name, _, story_name in activities:

        story = get_story(
            translated,
            story_name
        )

        if not story:
            continue

        add_heading(
            document,
            display_name,
            2
        )

        story_paragraph(
            document,
            story
        )


# ============================================================
# LEARNING OUTCOMES
# ============================================================

def add_learning_outcomes(
    document,
    translated
):

    add_heading(
        document,
        "Learning Outcomes",
        1
    )

    outcomes = translated.get(
        "learning_outcomes",
        {}
    )

    for name, title in [
        (
            "language",
            "Language Learning"
        ),
        (
            "mathematics",
            "Mathematics Learning"
        ),
    ]:

        section = outcomes.get(
            name,
            {}
        )

        if not section:
            continue

        add_heading(
            document,
            title,
            2
        )

        story = section.get(
            "story",
            ""
        )

        if story:

            add_body(
                document,
                story
            )

        add_bullets(
            document,
            section.get(
                "observed_progress",
                []
            )
        )


# ============================================================
# PARENT ENGAGEMENT
# ============================================================

def add_parent_engagement(
    document,
    translated
):

    section = translated.get(
        "parent_engagement",
        {}
    )

    if not section:
        return

    add_heading(
        document,
        "Parent and Community Engagement",
        1
    )

    add_body(
        document,
        "The programme works with parents and community "
        "stakeholders to strengthen children's regular "
        "participation and continued learning."
    )

    add_bullets(
        document,
        section.get(
            "key_observations",
            []
        )
    )

    examples = section.get(
        "examples",
        []
    )

    if examples:

        add_body(
            document,
            "Examples of engagement:"
        )

        add_bullets(
            document,
            examples
        )

    changes = section.get(
        "changes_observed",
        []
    )

    if changes:

        add_body(
            document,
            "Changes observed:"
        )

        add_bullets(
            document,
            changes
        )


# ============================================================
# TEACHER TRAINING
# ============================================================

def add_teacher_training(
    document,
    translated
):

    section = translated.get(
        "teacher_training",
        {}
    )

    if not section:
        return

    add_heading(
        document,
        "Teacher and Supervisor Capacity Building",
        1
    )

    add_body(
        document,
        "Regular capacity building supports teachers and "
        "supervisors in planning engaging sessions and "
        "responding to different learning needs."
    )

    topics = section.get(
        "topics",
        []
    )

    if topics:

        add_body(
            document,
            "Training areas included:"
        )

        add_bullets(
            document,
            topics
        )

    changes = section.get(
        "implementation_changes",
        []
    )

    if changes:

        add_body(
            document,
            "Changes in implementation:"
        )

        add_bullets(
            document,
            changes
        )


# ============================================================
# CHALLENGES
# ============================================================

def add_challenges(
    document,
    translated
):

    challenges = translated.get(
        "challenges",
        []
    )

    if not challenges:
        return

    add_heading(
        document,
        "Challenges",
        1
    )

    add_bullets(
        document,
        challenges
    )


# ============================================================
# FUTURE PLANNING
# ============================================================

def add_future_planning(
    document,
    translated
):

    plans = translated.get(
        "future_planning",
        []
    )

    if not plans:
        return

    add_heading(
        document,
        "Plan for the Next Period",
        1
    )

    add_bullets(
        document,
        plans
    )


# ============================================================
# SUCCESS STORIES
# ============================================================

def add_success_stories(
    document,
    translated
):

    stories = translated.get(
        "success_stories",
        []
    )

    if not stories:
        return

    add_heading(
        document,
        "Success Story",
        1
    )

    for story in stories:

        title = story.get(
            "title",
            "Success Story"
        )

        narrative = story.get(
            "story",
            ""
        )

        outcome = story.get(
            "outcome",
            ""
        )

        add_heading(
            document,
            title,
            2
        )

        if narrative:

            add_body(
                document,
                narrative
            )

        if outcome:

            add_body(
                document,
                f"Outcome: {outcome}"
            )


# ============================================================
# FOOTER
# ============================================================

def add_footer(
    document,
    funder
):

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            f"Door Step School Foundation | "
            f"Community Education Programme | "
            f"{FUNDER_NAMES.get(funder, funder)}"
        )

        run.font.size = Pt(8)


# ============================================================
# GENERATE REPORT
# ============================================================

def generate_report(
    funder,
    kpis,
    insights,
    translated
):

    document = Document()

    # Normal font
    styles = document.styles

    styles["Normal"].font.name = (
        "Arial"
    )

    styles["Normal"].font.size = Pt(
        10.5
    )

    styles["Heading 1"].font.name = (
        "Arial"
    )

    styles["Heading 1"].font.size = Pt(
        16
    )

    styles["Heading 1"].font.bold = True

    styles["Heading 2"].font.name = (
        "Arial"
    )

    styles["Heading 2"].font.size = Pt(
        13
    )

    styles["Heading 2"].font.bold = True

    # --------------------------------------------------------
    # Cover
    # --------------------------------------------------------

    create_cover(
        document,
        funder
    )

    # --------------------------------------------------------
    # First two pages / summary structure
    # --------------------------------------------------------

    add_introduction(
        document,
        translated
    )

    add_executive_summary(
        document,
        funder,
        kpis,
        translated
    )

    # --------------------------------------------------------
    # Main report
    # --------------------------------------------------------

    add_highlights(
        document,
        kpis,
        translated
    )

    add_activity_section(
        document,
        kpis,
        translated
    )

    add_learning_outcomes(
        document,
        translated
    )

    add_parent_engagement(
        document,
        translated
    )

    add_teacher_training(
        document,
        translated
    )

    add_challenges(
        document,
        translated
    )

    add_future_planning(
        document,
        translated
    )

    add_success_stories(
        document,
        translated
    )

    add_footer(
        document,
        funder
    )

    output_file = (
        OUTPUT_DIR /
        f"{funder}_Progress_Report_V4.docx"
    )

    document.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 65)
    print(
        "FUNDER-SPECIFIC CSR REPORT GENERATOR V4"
    )
    print("=" * 65)

    translated = load_translated_report()

    print(
        "\nLoaded:"
    )

    print(
        "output\\translated_report.json"
    )

    funders = []

    for file in KPI_DIR.glob(
        "*_kpis.json"
    ):

        funder = file.stem.replace(
            "_kpis",
            ""
        )

        if funder in FUNDER_NAMES:

            funders.append(
                funder
            )

    if not funders:

        raise RuntimeError(
            "No funder KPI files found."
        )

    print(
        f"\nFunders found: {len(funders)}"
    )

    for funder in sorted(
        funders
    ):

        print(
            f"\nGenerating report for: "
            f"{funder}"
        )

        kpis, insights = (
            load_funder_data(
                funder
            )
        )

        output_file = generate_report(
            funder,
            kpis,
            insights,
            translated
        )

        print(
            f"Created: {output_file}"
        )

    print("\n" + "=" * 65)

    print(
        "ALL V4 REPORTS GENERATED"
    )

    print("=" * 65)

    print(
        f"\nOutput folder:\n"
        f"{OUTPUT_DIR}"
    )


if __name__ == "__main__":
    main()