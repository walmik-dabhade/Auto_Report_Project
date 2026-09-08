import json
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

KPI_FILE = os.path.join(
    BASE_DIR,
    "output",
    "kpi_results.json"
)

INSIGHT_FILE = os.path.join(
    BASE_DIR,
    "output",
    "final_insights.json"
)

CONTENT_FILE = os.path.join(
    BASE_DIR,
    "report",
    "report_content.json"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output"
)

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "Automated_Report_V2.docx"
)


# ============================================================
# LOAD JSON
# ============================================================

def load_json(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_text(value):
    if value is None:
        return ""

    if isinstance(value, float):
        return str(round(value, 2))

    return str(value)


def find_value(data, possible_keys, default=0):
    """
    Searches recursively for a key.
    Makes the report generator more tolerant of changes
    in the KPI JSON structure.
    """

    if isinstance(data, dict):

        for key in possible_keys:
            if key in data:
                return data[key]

        for value in data.values():
            result = find_value(value, possible_keys, None)

            if result is not None:
                return result

    elif isinstance(data, list):

        for item in data:
            result = find_value(item, possible_keys, None)

            if result is not None:
                return result

    return default


def percentage(value):
    try:
        return f"{float(value):.2f}%"
    except:
        return "0.00%"


def add_page_number(paragraph):

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")

    run._r.append(field)


# ============================================================
# WORD FORMATTING
# ============================================================

def set_cell_text(cell, text, bold=False):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(clean_text(text))

    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):

    tcPr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(qn("w:fill"), fill)

    tcPr.append(shd)


def set_table_borders(table):

    tbl = table._tbl

    tblPr = tbl.tblPr

    borders = OxmlElement("w:tblBorders")

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):

        tag = OxmlElement(f"w:{edge}")

        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "B7B7B7")

        borders.append(tag)

    tblPr.append(borders)


def add_heading(doc, text, level=1):

    heading = doc.add_heading(text, level=level)

    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(6)

    return heading


def add_paragraph(doc, text, bold_start=None):

    paragraph = doc.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15

    if bold_start and text.startswith(bold_start):

        run1 = paragraph.add_run(bold_start)
        run1.bold = True

        run2 = paragraph.add_run(text[len(bold_start):])

    else:

        paragraph.add_run(clean_text(text))

    return paragraph


def add_bullet(doc, text):

    paragraph = doc.add_paragraph(
        style="List Bullet"
    )

    paragraph.paragraph_format.space_after = Pt(4)

    paragraph.add_run(clean_text(text))

    return paragraph


def add_numbered(doc, text):

    paragraph = doc.add_paragraph(
        style="List Number"
    )

    paragraph.paragraph_format.space_after = Pt(4)

    paragraph.add_run(clean_text(text))

    return paragraph


# ============================================================
# KPI EXTRACTION
# ============================================================

def get_overall_kpis(kpi):

    total = find_value(
        kpi,
        ["total_records", "total_rows", "total_children", "total"]
    )

    unique = find_value(
        kpi,
        [
            "unique_valid_software_ids",
            "unique_children",
            "unique_ids"
        ]
    )

    male = find_value(
        kpi,
        ["male", "male_children"]
    )

    female = find_value(
        kpi,
        ["female", "female_children"]
    )

    continued = find_value(
        kpi,
        ["continued"]
    )

    dropout = find_value(
        kpi,
        ["dropout", "dropouts"]
    )

    migrated = find_value(
        kpi,
        ["migrated", "migration"]
    )

    graduate = find_value(
        kpi,
        ["graduate", "graduates"]
    )

    continuation_rate = find_value(
        kpi,
        ["continuation_rate"]
    )

    dropout_rate = find_value(
        kpi,
        ["dropout_rate"]
    )

    migration_rate = find_value(
        kpi,
        ["migration_rate"]
    )

    graduate_rate = find_value(
        kpi,
        ["graduate_rate"]
    )

    return {
        "total": total,
        "unique": unique,
        "male": male,
        "female": female,
        "continued": continued,
        "dropout": dropout,
        "migrated": migrated,
        "graduate": graduate,
        "continuation_rate": continuation_rate,
        "dropout_rate": dropout_rate,
        "migration_rate": migration_rate,
        "graduate_rate": graduate_rate
    }


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def create_executive_summary(doc, kpis):

    add_heading(doc, "Executive Summary", 1)

    total = kpis["total"]
    unique = kpis["unique"]

    continued = kpis["continued"]
    dropout = kpis["dropout"]
    migrated = kpis["migrated"]
    graduate = kpis["graduate"]

    continuation_rate = kpis["continuation_rate"]
    dropout_rate = kpis["dropout_rate"]
    migration_rate = kpis["migration_rate"]

    text = (
        f"During the reporting period, the Community Education Program "
        f"recorded {total:,} child records representing {unique:,} unique "
        f"valid Software IDs. The programme continued to engage children "
        f"through educational activities focused on foundational learning, "
        f"regular participation and continued education."
    )

    add_paragraph(doc, text)

    text = (
        f"The latest available status data shows that {continued:,} children "
        f"({percentage(continuation_rate)}) are continuing, while "
        f"{dropout:,} children ({percentage(dropout_rate)}) are recorded "
        f"as dropouts and {migrated:,} children ({percentage(migration_rate)}) "
        f"are recorded as migrated. A further {graduate:,} children are "
        f"recorded as graduates."
    )

    add_paragraph(doc, text)

    if dropout > 0:

        add_paragraph(
            doc,
            f"The dropout figure indicates an important area for continued "
            f"follow-up. Understanding the reasons behind discontinuation and "
            f"strengthening engagement with children and parents can support "
            f"better continuity."
        )

    if migrated > 0:

        add_paragraph(
            doc,
            f"Migration is another important factor in continuity. Migrated "
            f"children should continue to be tracked wherever possible so "
            f"that movement between communities does not automatically "
            f"result in loss of educational support."
        )


# ============================================================
# STATIC CONTENT
# ============================================================

def add_static_content(doc, content):

    introduction = content.get(
        "introduction",
        {}
    )

    add_heading(
        doc,
        introduction.get(
            "title",
            "Introduction"
        ),
        1
    )

    for paragraph in introduction.get(
        "paragraphs",
        []
    ):

        add_paragraph(doc, paragraph)

    executive = content.get(
        "executive_summary",
        {}
    )

    add_heading(
        doc,
        "Programme Context",
        1
    )

    for paragraph in executive.get(
        "paragraphs",
        []
    ):

        add_paragraph(doc, paragraph)

    context = content.get(
        "programme_context",
        {}
    )

    for paragraph in context.get(
        "paragraphs",
        []
    ):

        add_paragraph(doc, paragraph)


# ============================================================
# KEY KPI TABLE
# ============================================================

def create_kpi_table(doc, kpis):

    add_heading(
        doc,
        "Key Programme Indicators",
        1
    )

    table = doc.add_table(
        rows=1,
        cols=3
    )

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    headers = [
        "Indicator",
        "Value",
        "Interpretation"
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            True
        )

        shade_cell(
            table.rows[0].cells[i],
            "D9EAF7"
        )

    rows = [

        (
            "Total Records",
            f"{kpis['total']:,}",
            "Total records in the cleaned dataset"
        ),

        (
            "Unique Valid Software IDs",
            f"{kpis['unique']:,}",
            "Unique valid child identifiers"
        ),

        (
            "Male Children",
            f"{kpis['male']:,}",
            "Male children recorded"
        ),

        (
            "Female Children",
            f"{kpis['female']:,}",
            "Female children recorded"
        ),

        (
            "Continued",
            f"{kpis['continued']:,}",
            f"{percentage(kpis['continuation_rate'])} continuation"
        ),

        (
            "Dropout",
            f"{kpis['dropout']:,}",
            f"{percentage(kpis['dropout_rate'])} dropout"
        ),

        (
            "Migrated",
            f"{kpis['migrated']:,}",
            f"{percentage(kpis['migration_rate'])} migration"
        ),

        (
            "Graduate",
            f"{kpis['graduate']:,}",
            "Children recorded as graduates"
        )
    ]

    for row in rows:

        cells = table.add_row().cells

        for i, value in enumerate(row):

            set_cell_text(
                cells[i],
                value
            )

    set_table_borders(table)


# ============================================================
# INSIGHT SECTIONS
# ============================================================

def add_insight_section(doc, title, insights):

    if not insights:
        return

    add_heading(
        doc,
        title,
        1
    )

    for insight in insights:

        add_numbered(
            doc,
            insight
        )


def get_insight_list(insights, key):

    value = insights.get(key, [])

    if isinstance(value, list):
        return value

    if isinstance(value, dict):

        for possible_key in [
            "insights",
            "items",
            "points"
        ]:

            if possible_key in value:

                if isinstance(
                    value[possible_key],
                    list
                ):
                    return value[possible_key]

    return []


# ============================================================
# ATTENDANCE SECTION
# ============================================================

def create_attendance_section(doc, kpi):

    attendance = kpi.get(
        "attendance",
        {}
    )

    if not attendance:
        return

    add_heading(
        doc,
        "Attendance",
        1
    )

    total = find_value(
        attendance,
        ["total"]
    )

    high = find_value(
        attendance,
        ["high_attendance", "high"]
    )

    medium = find_value(
        attendance,
        ["medium_attendance", "medium"]
    )

    low = find_value(
        attendance,
        ["low_attendance", "low"]
    )

    high_rate = find_value(
        attendance,
        ["high_attendance_rate", "high_rate"]
    )

    medium_rate = find_value(
        attendance,
        ["medium_attendance_rate", "medium_rate"]
    )

    low_rate = find_value(
        attendance,
        ["low_attendance_rate", "low_rate"]
    )

    add_paragraph(
        doc,
        f"Attendance data is available for {total:,} children. "
        f"{high:,} children ({percentage(high_rate)}) fall within "
        f"the high-attendance category, while {medium:,} children "
        f"({percentage(medium_rate)}) have medium attendance and "
        f"{low:,} children ({percentage(low_rate)}) have low attendance."
    )

    add_paragraph(
        doc,
        f"The attendance pattern indicates that high attendance is the "
        f"dominant category. At the same time, the {low:,} children in "
        f"the low-attendance group should remain a priority for follow-up "
        f"with teachers, parents and programme teams."
    )


# ============================================================
# LANGUAGE SECTION
# ============================================================

def create_language_section(doc, kpi):

    language = kpi.get(
        "language",
        {}
    )

    if not language:
        return

    add_heading(
        doc,
        "Language Learning",
        1
    )

    l1 = find_value(language, ["l1"])
    l2 = find_value(language, ["l2"])
    l3 = find_value(language, ["l3"])
    l4 = find_value(language, ["l4"])

    assessed = find_value(
        language,
        ["total_assessed", "assessed"]
    )

    missing = find_value(
        language,
        ["missing"]
    )

    lower = find_value(
        language,
        ["l1_l2", "l1_plus_l2"]
    )

    higher = find_value(
        language,
        ["l3_l4", "l3_plus_l4"]
    )

    lower_rate = find_value(
        language,
        ["l1_l2_rate", "l1_plus_l2_rate"]
    )

    higher_rate = find_value(
        language,
        ["l3_l4_rate", "l3_plus_l4_rate"]
    )

    add_paragraph(
        doc,
        f"Language assessment data is available for {assessed:,} children. "
        f"{missing:,} children do not have a recorded language assessment "
        f"in the latest available assessment data."
    )

    add_paragraph(
        doc,
        f"The higher language levels (L3 and L4) account for "
        f"{higher:,} assessed children ({percentage(higher_rate)}), "
        f"while L1 and L2 account for {lower:,} children "
        f"({percentage(lower_rate)})."
    )

    table = doc.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    headers = [
        "Language Level",
        "Children",
        "Share of Assessed"
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            True
        )

        shade_cell(
            table.rows[0].cells[i],
            "E2F0D9"
        )

    levels = [
        ("L1", l1),
        ("L2", l2),
        ("L3", l3),
        ("L4", l4)
    ]

    for level, count in levels:

        rate = (
            (count / assessed) * 100
            if assessed
            else 0
        )

        cells = table.add_row().cells

        set_cell_text(cells[0], level)
        set_cell_text(cells[1], f"{count:,}")
        set_cell_text(cells[2], percentage(rate))

    set_table_borders(table)


# ============================================================
# MATHEMATICS SECTION
# ============================================================

def create_mathematics_section(doc, kpi):

    maths = kpi.get(
        "mathematics",
        kpi.get("maths", {})
    )

    if not maths:
        return

    add_heading(
        doc,
        "Mathematics Learning",
        1
    )

    assessed = find_value(
        maths,
        ["total_assessed", "assessed"]
    )

    missing = find_value(
        maths,
        ["missing"]
    )

    low = find_value(
        maths,
        [
            "m1_m2_m3",
            "m1_plus_m2_plus_m3"
        ]
    )

    middle = find_value(
        maths,
        ["m4_m5"]
    )

    high = find_value(
        maths,
        ["m6_m7"]
    )

    low_rate = find_value(
        maths,
        [
            "m1_m2_m3_rate",
            "m1_plus_m2_plus_m3_rate"
        ]
    )

    middle_rate = find_value(
        maths,
        ["m4_m5_rate"]
    )

    high_rate = find_value(
        maths,
        ["m6_m7_rate"]
    )

    add_paragraph(
        doc,
        f"Mathematics assessment data is available for {assessed:,} "
        f"children. A total of {missing:,} children do not have a "
        f"recorded mathematics assessment in the latest available data."
    )

    add_paragraph(
        doc,
        f"Among assessed children, {high:,} ({percentage(high_rate)}) "
        f"are at M6 or M7, while {low:,} ({percentage(low_rate)}) "
        f"are at M1 to M3. A further {middle:,} children "
        f"({percentage(middle_rate)}) are at M4 or M5."
    )

    table = doc.add_table(
        rows=1,
        cols=3
    )

    table.style = "Table Grid"

    headers = [
        "Mathematics Level",
        "Children",
        "Share"
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            table.rows[0].cells[i],
            header,
            True
        )

        shade_cell(
            table.rows[0].cells[i],
            "FFF2CC"
        )

    for level in [
        "M1",
        "M2",
        "M3",
        "M4",
        "M5",
        "M6",
        "M7"
    ]:

        count = find_value(
            maths,
            [level.lower()]
        )

        rate = (
            (count / assessed) * 100
            if assessed
            else 0
        )

        cells = table.add_row().cells

        set_cell_text(cells[0], level)
        set_cell_text(cells[1], f"{count:,}")
        set_cell_text(cells[2], percentage(rate))

    set_table_borders(table)


# ============================================================
# MAIN REPORT
# ============================================================

def generate_report():

    print()
    print("========================================")
    print("AUTOMATED REPORT GENERATOR V2")
    print("========================================")
    print()

    print("Loading KPI data...")

    kpi = load_json(KPI_FILE)

    print("KPI data loaded successfully.")

    print("Loading insights...")

    insights = load_json(INSIGHT_FILE)

    print("Insights loaded successfully.")

    print("Loading report content...")

    content = load_json(CONTENT_FILE)

    print("Static report content loaded successfully.")

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    doc = Document()

    # --------------------------------------------------------
    # PAGE SETUP
    # --------------------------------------------------------

    section = doc.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # --------------------------------------------------------
    # DEFAULT FONT
    # --------------------------------------------------------

    styles = doc.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(24)

    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)

    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)

    # --------------------------------------------------------
    # COVER PAGE
    # --------------------------------------------------------

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    paragraph.space_after = Pt(30)

    run = paragraph.add_run(
        content.get(
            "organization_name",
            "Door Step School Foundation"
        )
    )

    run.bold = True
    run.font.size = Pt(20)

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        content.get(
            "program_name",
            "Community Education Program"
        )
    )

    run.bold = True
    run.font.size = Pt(25)

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "PROJECT PROGRESS REPORT"
    )

    run.bold = True
    run.font.size = Pt(18)

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "Automated Programme Performance Report"
    )

    run.font.size = Pt(13)

    paragraph = doc.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(
        "\nGenerated from the latest cleaned programme dataset"
    )

    run.font.size = Pt(10)

    doc.add_page_break()

    # --------------------------------------------------------
    # INTRODUCTION + STATIC STORY
    # --------------------------------------------------------

    add_static_content(
        doc,
        content
    )

    # --------------------------------------------------------
    # EXECUTIVE SUMMARY
    # --------------------------------------------------------

    kpis = get_overall_kpis(kpi)

    create_executive_summary(
        doc,
        kpis
    )

    # --------------------------------------------------------
    # KEY KPI TABLE
    # --------------------------------------------------------

    create_kpi_table(
        doc,
        kpis
    )

    # --------------------------------------------------------
    # OVERALL INSIGHTS
    # --------------------------------------------------------

    add_insight_section(
        doc,
        "Overall Insights",
        get_insight_list(
            insights,
            "overall"
        )
    )

    # --------------------------------------------------------
    # ACTIVITY INSIGHTS
    # --------------------------------------------------------

    add_insight_section(
        doc,
        "Activity Performance and Insights",
        get_insight_list(
            insights,
            "activity"
        )
    )

    # --------------------------------------------------------
    # FUNDER INSIGHTS
    # --------------------------------------------------------

    add_insight_section(
        doc,
        "Funder Performance and Insights",
        get_insight_list(
            insights,
            "funder"
        )
    )

    # --------------------------------------------------------
    # ATTENDANCE
    # --------------------------------------------------------

    create_attendance_section(
        doc,
        kpi
    )

    add_insight_section(
        doc,
        "Attendance Insights",
        get_insight_list(
            insights,
            "attendance"
        )
    )

    # --------------------------------------------------------
    # LANGUAGE
    # --------------------------------------------------------

    create_language_section(
        doc,
        kpi
    )

    add_insight_section(
        doc,
        "Language Learning Insights",
        get_insight_list(
            insights,
            "language"
        )
    )

    # --------------------------------------------------------
    # MATHEMATICS
    # --------------------------------------------------------

    create_mathematics_section(
        doc,
        kpi
    )

    add_insight_section(
        doc,
        "Mathematics Learning Insights",
        get_insight_list(
            insights,
            "mathematics"
        )
    )

    # --------------------------------------------------------
    # LOCATION
    # --------------------------------------------------------

    add_insight_section(
        doc,
        "Location Performance and Insights",
        get_insight_list(
            insights,
            "location"
        )
    )

    # --------------------------------------------------------
    # PROGRAMME STORY
    # --------------------------------------------------------

    add_heading(
        doc,
        "Programme Approach and Learning Journey",
        1
    )

    add_paragraph(
        doc,
        "The programme combines structured learning with engaging "
        "activities designed to build children's interest, improve "
        "participation and strengthen foundational literacy and "
        "numeracy. The approach also places importance on regular "
        "tracking, parent participation and follow-up with children "
        "who require additional support."
    )

    add_paragraph(
        doc,
        "The reference programme material highlights the importance "
        "of innovative and consistent teaching methods, foundational "
        "literacy and numeracy, children's overall development and "
        "parent participation in supporting continued education."
    )

    # --------------------------------------------------------
    # CHALLENGES
    # --------------------------------------------------------

    add_heading(
        doc,
        "Key Challenges and Areas Requiring Attention",
        1
    )

    dropout = kpis["dropout"]
    migrated = kpis["migrated"]

    if dropout > 0:

        add_bullet(
            doc,
            f"{dropout:,} children are recorded as dropouts in the "
            f"latest status data. Reasons for discontinuation should "
            f"be reviewed and categorised wherever possible."
        )

    if migrated > 0:

        add_bullet(
            doc,
            f"{migrated:,} children are recorded as migrated. "
            f"Tracking their educational continuity after migration "
            f"can help distinguish migration-related movement from "
            f"actual disengagement."
        )

    language = kpi.get("language", {})

    language_missing = find_value(
        language,
        ["missing"]
    )

    if language_missing > 0:

        add_bullet(
            doc,
            f"{language_missing:,} children do not have a recorded "
            f"language assessment in the latest available data."
        )

    maths = kpi.get(
        "mathematics",
        kpi.get("maths", {})
    )

    maths_missing = find_value(
        maths,
        ["missing"]
    )

    if maths_missing > 0:

        add_bullet(
            doc,
            f"{maths_missing:,} children do not have a recorded "
            f"mathematics assessment in the latest available data."
        )

    # --------------------------------------------------------
    # RECOMMENDATIONS
    # --------------------------------------------------------

    add_heading(
        doc,
        "Recommendations",
        1
    )

    recommendations = [

        "Prioritise follow-up with children showing low attendance "
        "and identify the specific reasons affecting their regular participation.",

        "Review dropout cases individually and classify the reasons "
        "for discontinuation so that programme responses can be better targeted.",

        "Continue tracking migrated children and document whether "
        "they continue education after moving to another location.",

        "Strengthen assessment coverage for language and mathematics "
        "so that learning progress can be monitored more consistently.",

        "Use activity-level and location-level performance to identify "
        "areas requiring additional support and replicate approaches "
        "from locations or activities showing stronger continuation.",

        "Continue parent engagement through regular communication, "
        "meetings and practical guidance on supporting children's learning at home."
    ]

    for recommendation in recommendations:

        add_numbered(
            doc,
            recommendation
        )

    # --------------------------------------------------------
    # STATIC REPORT NOTE
    # --------------------------------------------------------

    report_note = content.get(
        "report_note",
        {}
    )

    if report_note:

        add_heading(
            doc,
            report_note.get(
                "title",
                "About This Report"
            ),
            1
        )

        add_paragraph(
            doc,
            report_note.get(
                "text",
                ""
            )
        )

    # --------------------------------------------------------
    # FOOTERS
    # --------------------------------------------------------

    for section in doc.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run(
            "Community Education Program | Automated Progress Report | "
        )

        add_page_number(
            paragraph
        )

    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    doc.save(
        OUTPUT_FILE
    )

    print()
    print("========================================")
    print("REPORT GENERATION COMPLETED")
    print("========================================")
    print()
    print("Report saved to:")
    print(OUTPUT_FILE)
    print()


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    generate_report()