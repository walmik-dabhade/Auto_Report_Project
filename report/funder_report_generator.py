import json
import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

KPI_DIR = os.path.join(
    BASE_DIR,
    "output",
    "funder_kpis"
)

INSIGHT_DIR = os.path.join(
    BASE_DIR,
    "output",
    "funder_insights"
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output",
    "funder_reports"
)

os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)


# ============================================================
# REPORT INFORMATION
# ============================================================

REPORT_PERIOD = "April 2025 – December 2025"

ORGANISATION_NAME = (
    "DOOR STEP SCHOOL FOUNDATION"
)

PROGRAMME_NAME = (
    "COMMUNITY EDUCATION PROGRAMME"
)


# ============================================================
# GENERAL FORMATTING
# ============================================================

def set_cell_shading(cell, fill):

    tcPr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(
        qn("w:fill"),
        fill
    )

    tcPr.append(shd)


def set_cell_text(
    cell,
    text,
    bold=False,
    size=9
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(
        str(text)
    )

    run.bold = bold

    run.font.size = Pt(size)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )


def add_heading(
    document,
    text,
    level=1
):

    heading = document.add_heading(
        text,
        level=level
    )

    return heading


def add_paragraph(
    document,
    text,
    bold_start=None
):

    paragraph = document.add_paragraph()

    if bold_start and text.startswith(
        bold_start
    ):

        first = paragraph.add_run(
            bold_start
        )

        first.bold = True

        paragraph.add_run(
            text[len(bold_start):]
        )

    else:

        paragraph.add_run(
            text
        )

    return paragraph


def add_bullet(
    document,
    text
):

    paragraph = document.add_paragraph(
        style="List Bullet"
    )

    paragraph.add_run(
        text
    )

    return paragraph


def add_numbered(
    document,
    number,
    text
):

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        f"{number}. "
    )

    run.bold = True

    paragraph.add_run(
        text
    )

    return paragraph


# ============================================================
# TABLE CREATION
# ============================================================

def create_table(
    document,
    headers,
    rows
):

    table = document.add_table(
        rows=1,
        cols=len(headers)
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    for index, header in enumerate(
        headers
    ):

        set_cell_text(
            header_cells[index],
            header,
            bold=True,
            size=9
        )

        set_cell_shading(
            header_cells[index],
            "D9EAF7"
        )

    for row in rows:

        cells = table.add_row().cells

        for index, value in enumerate(
            row
        ):

            set_cell_text(
                cells[index],
                value,
                size=9
            )

    document.add_paragraph()

    return table


# ============================================================
# COVER PAGE
# ============================================================

def create_cover(
    document,
    funder,
    kpi
):

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.add_paragraph()
    document.add_paragraph()
    document.add_paragraph()

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = title.add_run(
        ORGANISATION_NAME
    )

    run.bold = True
    run.font.size = Pt(24)

    subtitle = document.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = subtitle.add_run(
        PROGRAMME_NAME
    )

    run.bold = True
    run.font.size = Pt(18)

    document.add_paragraph()

    report_title = document.add_paragraph()

    report_title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = report_title.add_run(
        "FUNDER-SPECIFIC PROGRESS REPORT"
    )

    run.bold = True
    run.font.size = Pt(20)

    document.add_paragraph()

    funder_title = document.add_paragraph()

    funder_title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = funder_title.add_run(
        funder
    )

    run.bold = True
    run.font.size = Pt(26)

    period = document.add_paragraph()

    period.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = period.add_run(
        REPORT_PERIOD
    )

    run.font.size = Pt(14)

    document.add_paragraph()
    document.add_paragraph()

    overall = kpi.get(
        "overall",
        {}
    )

    status = kpi.get(
        "status",
        {}
    )

    summary_rows = [

        [
            "Children / Records",
            f"{overall.get('total_records', 0):,}"
        ],

        [
            "Unique Valid IDs",
            f"{overall.get('unique_valid_software_ids', 0):,}"
        ],

        [
            "Continued",
            f"{status.get('continued', 0):,}"
        ],

        [
            "Continuation Rate",
            f"{status.get('continuation_rate', 0):.2f}%"
        ],

        [
            "Dropout Rate",
            f"{status.get('dropout_rate', 0):.2f}%"
        ],

        [
            "Migration Rate",
            f"{status.get('migration_rate', 0):.2f}%"
        ]

    ]

    table = create_table(
        document,
        ["Key Indicator", "Value"],
        summary_rows
    )

    document.add_page_break()


# ============================================================
# PROGRAMME INTRODUCTION
# ============================================================

def add_programme_introduction(
    document
):

    add_heading(
        document,
        "1. Programme Overview",
        1
    )

    add_paragraph(
        document,
        "The Community Education Programme works with "
        "children from vulnerable communities who may face "
        "barriers to regular educational participation. "
        "The programme focuses on improving access to "
        "education, supporting continuity and strengthening "
        "foundational learning."
    )

    add_paragraph(
        document,
        "The programme uses community-based educational "
        "activities to reach children in their local "
        "environments. The approach combines learning "
        "support, reading opportunities, structured "
        "activities and continued engagement with children "
        "and families."
    )

    add_paragraph(
        document,
        "The purpose of this report is to present the "
        "progress of the programme supported by the "
        "respective funder during the reporting period, "
        "combining quantitative programme data with "
        "interpretation of the observed results."
    )


# ============================================================
# EXECUTIVE SUMMARY
# ============================================================

def add_executive_summary(
    document,
    funder,
    kpi
):

    add_heading(
        document,
        "2. Executive Summary",
        1
    )

    overall = kpi.get(
        "overall",
        {}
    )

    status = kpi.get(
        "status",
        {}
    )

    total = overall.get(
        "total_records",
        0
    )

    unique = overall.get(
        "unique_valid_software_ids",
        0
    )

    continued = status.get(
        "continued",
        0
    )

    continuation_rate = status.get(
        "continuation_rate",
        0
    )

    dropout = status.get(
        "dropout",
        0
    )

    dropout_rate = status.get(
        "dropout_rate",
        0
    )

    migrated = status.get(
        "migrated",
        0
    )

    migration_rate = status.get(
        "migration_rate",
        0
    )

    add_paragraph(
        document,
        f"During {REPORT_PERIOD}, {funder} supported "
        f"{total:,} child records representing "
        f"{unique:,} unique valid Software IDs."
    )

    add_paragraph(
        document,
        f"The latest status data records "
        f"{continued:,} children ({continuation_rate:.2f}%) "
        f"as continuing with the programme. "
        f"{dropout:,} children ({dropout_rate:.2f}%) "
        f"are recorded as dropouts, while "
        f"{migrated:,} children ({migration_rate:.2f}%) "
        f"are recorded as migrated."
    )

    attendance = kpi.get(
        "attendance",
        {}
    )

    if attendance:

        high = attendance.get(
            "high_attendance",
            0
        )

        high_rate = attendance.get(
            "high_attendance_rate",
            0
        )

        add_paragraph(
            document,
            f"Attendance data shows that {high:,} children "
            f"({high_rate:.2f}%) fall within the high "
            f"attendance category of 80–100%."
        )

    language = kpi.get(
        "language",
        {}
    )

    if language:

        assessed = language.get(
            "total_assessed",
            0
        )

        higher = language.get(
            "l3_l4",
            0
        )

        higher_rate = language.get(
            "l3_l4_rate",
            0
        )

        if assessed:

            add_paragraph(
                document,
                f"Among the {assessed:,} children with "
                f"recorded language assessments, "
                f"{higher:,} ({higher_rate:.2f}%) "
                f"are at L3 or L4."
            )

    mathematics = kpi.get(
        "mathematics",
        {}
    )

    if mathematics:

        assessed = mathematics.get(
            "total_assessed",
            0
        )

        higher = mathematics.get(
            "m6_m7",
            0
        )

        higher_rate = mathematics.get(
            "m6_m7_rate",
            0
        )

        if assessed:

            add_paragraph(
                document,
                f"Among the {assessed:,} children with "
                f"recorded mathematics assessments, "
                f"{higher:,} ({higher_rate:.2f}%) "
                f"are at M6 or M7."
            )


# ============================================================
# KEY HIGHLIGHTS
# ============================================================

def add_key_highlights(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "3. Key Highlights",
        1
    )

    status = kpi.get(
        "status",
        {}
    )

    attendance = kpi.get(
        "attendance",
        {}
    )

    language = kpi.get(
        "language",
        {}
    )

    mathematics = kpi.get(
        "mathematics",
        {}
    )

    highlights = []

    if status:

        highlights.append(
            f"Continuation: "
            f"{status.get('continued', 0):,} children "
            f"({status.get('continuation_rate', 0):.2f}%)."
        )

    if attendance:

        highlights.append(
            f"High attendance: "
            f"{attendance.get('high_attendance', 0):,} children "
            f"({attendance.get('high_attendance_rate', 0):.2f}%)."
        )

    if language:

        highlights.append(
            f"Language L3/L4: "
            f"{language.get('l3_l4', 0):,} assessed children "
            f"({language.get('l3_l4_rate', 0):.2f}%)."
        )

    if mathematics:

        highlights.append(
            f"Mathematics M6/M7: "
            f"{mathematics.get('m6_m7', 0):,} assessed children "
            f"({mathematics.get('m6_m7_rate', 0):.2f}%)."
        )

    for item in highlights:

        add_bullet(
            document,
            item
        )

    document.add_paragraph()

    # Add selected insights
    overall_insights = insights.get(
        "overall",
        []
    )

    for insight in overall_insights[:3]:

        add_paragraph(
            document,
            insight
        )


# ============================================================
# GENDER
# ============================================================

def add_gender_section(
    document,
    kpi
):

    add_heading(
        document,
        "4. Gender Profile",
        1
    )

    overall = kpi.get(
        "overall",
        {}
    )

    male = overall.get(
        "male",
        0
    )

    female = overall.get(
        "female",
        0
    )

    total = overall.get(
        "total_records",
        0
    )

    male_rate = (
        male / total * 100
        if total else 0
    )

    female_rate = (
        female / total * 100
        if total else 0
    )

    create_table(
        document,

        [
            "Gender",
            "Children",
            "Share"
        ],

        [
            [
                "Male",
                f"{male:,}",
                f"{male_rate:.2f}%"
            ],

            [
                "Female",
                f"{female:,}",
                f"{female_rate:.2f}%"
            ]
        ]
    )

    if female > male:

        add_paragraph(
            document,
            f"Female children represent the larger share "
            f"of the supported population, accounting for "
            f"{female_rate:.2f}% of records."
        )

    elif male > female:

        add_paragraph(
            document,
            f"Male children represent the larger share "
            f"of the supported population, accounting for "
            f"{male_rate:.2f}% of records."
        )


# ============================================================
# STATUS
# ============================================================

def add_status_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "5. Programme Continuity and Status",
        1
    )

    status = kpi.get(
        "status",
        {}
    )

    rows = [

        [
            "Continued",
            status.get("continued", 0),
            f"{status.get('continuation_rate', 0):.2f}%"
        ],

        [
            "Dropout",
            status.get("dropout", 0),
            f"{status.get('dropout_rate', 0):.2f}%"
        ],

        [
            "Migrated",
            status.get("migrated", 0),
            f"{status.get('migration_rate', 0):.2f}%"
        ],

        [
            "Graduate",
            status.get("graduate", 0),
            f"{status.get('graduate_rate', 0):.2f}%"
        ]

    ]

    create_table(
        document,
        [
            "Status",
            "Children",
            "Rate"
        ],
        rows
    )

    add_paragraph(
        document,
        "The continuity figures provide an indication of "
        "how effectively children remain engaged with the "
        "programme. Dropout and migration should be "
        "interpreted separately because migration may reflect "
        "movement of families rather than educational "
        "disengagement."
    )

    for insight in insights.get(
        "overall",
        []
    ):

        if (
            "continued" in insight.lower()
            or "dropout" in insight.lower()
            or "migrated" in insight.lower()
        ):

            add_paragraph(
                document,
                insight
            )


# ============================================================
# ACTIVITY
# ============================================================

def add_activity_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "6. Activity-wise Programme Performance",
        1
    )

    activities = kpi.get(
        "activity",
        {}
    )

    if not activities:

        add_paragraph(
            document,
            "No activity-level data is available."
        )

        return

    rows = []

    for activity, values in activities.items():

        rows.append(

            [
                activity,

                f"{values.get('total', 0):,}",

                f"{values.get('continued', 0):,}",

                f"{values.get('continuation_rate', 0):.2f}%",

                f"{values.get('dropout_rate', 0):.2f}%",

                f"{values.get('migration_rate', 0):.2f}%"

            ]

        )

    create_table(
        document,

        [
            "Activity",
            "Total",
            "Continued",
            "Continuation",
            "Dropout",
            "Migration"
        ],

        rows
    )

    add_paragraph(
        document,
        "Activity-level differences help identify where "
        "programme engagement is particularly strong and "
        "where additional review may be required."
    )

    for insight in insights.get(
        "activity",
        []
    ):

        add_bullet(
            document,
            insight
        )


# ============================================================
# ATTENDANCE
# ============================================================

def add_attendance_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "7. Attendance Performance",
        1
    )

    attendance = kpi.get(
        "attendance",
        {}
    )

    if not attendance:

        add_paragraph(
            document,
            "Attendance data is not available."
        )

        return

    rows = [

        [
            "High (80–100%)",
            attendance.get(
                "high_attendance",
                0
            ),
            f"{attendance.get('high_attendance_rate', 0):.2f}%"
        ],

        [
            "Medium (51–79%)",
            attendance.get(
                "medium_attendance",
                0
            ),
            f"{attendance.get('medium_attendance_rate', 0):.2f}%"
        ],

        [
            "Low (1–50%)",
            attendance.get(
                "low_attendance",
                0
            ),
            f"{attendance.get('low_attendance_rate', 0):.2f}%"
        ]

    ]

    create_table(
        document,
        [
            "Attendance Category",
            "Children",
            "Share"
        ],
        rows
    )

    add_paragraph(
        document,
        "Regular attendance is an important indicator of "
        "continued engagement. Children in the low-attendance "
        "category may benefit from closer follow-up with "
        "families and programme teams."
    )

    for insight in insights.get(
        "attendance",
        []
    ):

        add_bullet(
            document,
            insight
        )


# ============================================================
# LANGUAGE
# ============================================================

def add_language_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "8. Language Learning Outcomes",
        1
    )

    language = kpi.get(
        "language",
        {}
    )

    if not language:

        add_paragraph(
            document,
            "Language assessment data is not available."
        )

        return

    assessed = language.get(
        "total_assessed",
        0
    )

    missing = language.get(
        "missing",
        0
    )

    rows = []

    for level in [
        "L1",
        "L2",
        "L3",
        "L4"
    ]:

        key = level.lower()

        count = language.get(
            key,
            0
        )

        rate = (
            count / assessed * 100
            if assessed else 0
        )

        rows.append(
            [
                level,
                f"{count:,}",
                f"{rate:.2f}%"
            ]
        )

    create_table(
        document,
        [
            "Language Level",
            "Children",
            "Share of Assessed"
        ],
        rows
    )

    add_paragraph(
        document,
        f"Language assessment data is available for "
        f"{assessed:,} children. {missing:,} children do not "
        f"have a recorded language assessment in the latest "
        f"assessment column."
    )

    higher = language.get(
        "l3_l4",
        0
    )

    higher_rate = language.get(
        "l3_l4_rate",
        0
    )

    lower = language.get(
        "l1_l2",
        0
    )

    lower_rate = language.get(
        "l1_l2_rate",
        0
    )

    add_paragraph(
        document,
        f"Among assessed children, {higher:,} "
        f"({higher_rate:.2f}%) are at L3 or L4, while "
        f"{lower:,} ({lower_rate:.2f}%) are at L1 or L2."
    )

    for insight in insights.get(
        "language",
        []
    ):

        add_bullet(
            document,
            insight
        )


# ============================================================
# MATHEMATICS
# ============================================================

def add_mathematics_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "9. Mathematics Learning Outcomes",
        1
    )

    mathematics = kpi.get(
        "mathematics",
        {}
    )

    if not mathematics:

        add_paragraph(
            document,
            "Mathematics assessment data is not available."
        )

        return

    assessed = mathematics.get(
        "total_assessed",
        0
    )

    missing = mathematics.get(
        "missing",
        0
    )

    rows = []

    for level in [
        "M1",
        "M2",
        "M3",
        "M4",
        "M5",
        "M6",
        "M7"
    ]:

        key = level.lower()

        count = mathematics.get(
            key,
            0
        )

        rate = (
            count / assessed * 100
            if assessed else 0
        )

        rows.append(
            [
                level,
                f"{count:,}",
                f"{rate:.2f}%"
            ]
        )

    create_table(
        document,
        [
            "Mathematics Level",
            "Children",
            "Share of Assessed"
        ],
        rows
    )

    add_paragraph(
        document,
        f"Mathematics assessment data is available for "
        f"{assessed:,} children. {missing:,} children do not "
        f"have a recorded mathematics assessment in the latest "
        f"assessment column."
    )

    low = mathematics.get(
        "m1_m2_m3",
        0
    )

    low_rate = mathematics.get(
        "m1_m2_m3_rate",
        0
    )

    middle = mathematics.get(
        "m4_m5",
        0
    )

    middle_rate = mathematics.get(
        "m4_m5_rate",
        0
    )

    higher = mathematics.get(
        "m6_m7",
        0
    )

    higher_rate = mathematics.get(
        "m6_m7_rate",
        0
    )

    add_paragraph(
        document,
        f"Lower levels (M1–M3) account for {low:,} "
        f"children ({low_rate:.2f}%), middle levels (M4–M5) "
        f"account for {middle:,} ({middle_rate:.2f}%), while "
        f"higher levels (M6–M7) account for {higher:,} "
        f"({higher_rate:.2f}%)."
    )

    for insight in insights.get(
        "mathematics",
        []
    ):

        add_bullet(
            document,
            insight
        )


# ============================================================
# LOCATION
# ============================================================

def add_location_section(
    document,
    kpi,
    insights
):

    add_heading(
        document,
        "10. Location-wise Performance",
        1
    )

    location = kpi.get(
        "location",
        {}
    )

    locations = location.get(
        "locations",
        {}
    )

    if locations:

        rows = []

        for name, values in locations.items():

            rows.append(

                [
                    name,

                    f"{values.get('total', 0):,}",

                    f"{values.get('continued', 0):,}",

                    f"{values.get('continuation_rate', 0):.2f}%",

                    f"{values.get('dropout_rate', 0):.2f}%",

                    f"{values.get('migration_rate', 0):.2f}%"

                ]

            )

        create_table(
            document,

            [
                "Location",
                "Total",
                "Continued",
                "Continuation",
                "Dropout",
                "Migration"
            ],

            rows
        )

    areas = location.get(
        "areas",
        {}
    )

    if areas:

        add_heading(
            document,
            "Area Distribution",
            2
        )

        total = sum(
            areas.values()
        )

        rows = []

        for area, count in areas.items():

            rate = (
                count / total * 100
                if total else 0
            )

            rows.append(
                [
                    area,
                    f"{count:,}",
                    f"{rate:.2f}%"
                ]
            )

        create_table(
            document,
            [
                "Area",
                "Children",
                "Share"
            ],
            rows
        )

    for insight in insights.get(
        "location",
        []
    ):

        add_bullet(
            document,
            insight
        )


# ============================================================
# PROGRAMME INTERPRETATION
# ============================================================

def add_interpretation(
    document,
    funder,
    kpi,
    insights
):

    add_heading(
        document,
        "11. Interpretation of Programme Performance",
        1
    )

    status = kpi.get(
        "status",
        {}
    )

    continuation = status.get(
        "continuation_rate",
        0
    )

    dropout = status.get(
        "dropout_rate",
        0
    )

    migration = status.get(
        "migration_rate",
        0
    )

    if continuation >= 90:

        continuity_text = (
            "The continuation rate is very strong, "
            "indicating sustained engagement among "
            "the children supported by the programme."
        )

    elif continuation >= 80:

        continuity_text = (
            "The continuation rate indicates that a "
            "large majority of supported children remain "
            "engaged, while continued attention to dropout "
            "and migration can further strengthen continuity."
        )

    else:

        continuity_text = (
            "The continuation rate indicates that "
            "programme continuity requires particular "
            "attention. Understanding the reasons behind "
            "dropout and migration should be an important "
            "management priority."
        )

    add_paragraph(
        document,
        f"For {funder}, the latest programme data indicates "
        f"an overall continuation rate of {continuation:.2f}%. "
        f"{continuity_text}"
    )

    if dropout >= 15:

        add_paragraph(
            document,
            f"The dropout rate of {dropout:.2f}% is relatively "
            f"high and warrants focused review of the activities "
            f"and locations where dropout is concentrated."
        )

    elif dropout > 0:

        add_paragraph(
            document,
            f"The recorded dropout rate is {dropout:.2f}%. "
            f"Continued monitoring and timely follow-up can "
            f"help identify emerging continuity risks."
        )

    if migration >= 15:

        add_paragraph(
            document,
            f"The migration rate of {migration:.2f}% is notable. "
            f"Because migration may reflect movement of families, "
            f"tracking relocated children can help distinguish "
            f"mobility from programme disengagement."
        )

    attendance = kpi.get(
        "attendance",
        {}
    )

    if attendance:

        low_rate = attendance.get(
            "low_attendance_rate",
            0
        )

        if low_rate >= 10:

            add_paragraph(
                document,
                f"Low attendance affects {low_rate:.2f}% of "
                f"children with attendance records, suggesting "
                f"that targeted attendance follow-up could "
                f"support stronger continuity."
            )


# ============================================================
# PRIORITY ACTIONS
# ============================================================

def add_recommendations(
    document,
    funder,
    kpi,
    insights
):

    add_heading(
        document,
        "12. Priority Areas and Recommended Actions",
        1
    )

    status = kpi.get(
        "status",
        {}
    )

    attendance = kpi.get(
        "attendance",
        {}
    )

    language = kpi.get(
        "language",
        {}
    )

    mathematics = kpi.get(
        "mathematics",
        {}
    )

    actions = []

    if status.get(
        "dropout_rate",
        0
    ) >= 10:

        actions.append(
            "Review dropout patterns at activity and "
            "location level and conduct targeted follow-up "
            "for children at risk of disengagement."
        )

    if status.get(
        "migration_rate",
        0
    ) >= 10:

        actions.append(
            "Strengthen migration tracking so that children "
            "who relocate can be followed where feasible "
            "and distinguished from programme dropouts."
        )

    if attendance:

        if attendance.get(
            "low_attendance_rate",
            0
        ) >= 10:

            actions.append(
                "Prioritise children with low attendance "
                "for family engagement and follow-up."
            )

    if language:

        if language.get(
            "missing",
            0
        ) > 0:

            actions.append(
                "Improve language assessment coverage so "
                "that learning progress can be tracked "
                "for a larger share of supported children."
            )

    if mathematics:

        if mathematics.get(
            "missing",
            0
        ) > 0:

            actions.append(
                "Strengthen mathematics assessment coverage "
                "and use lower-level results to identify "
                "children requiring additional learning support."
            )

    activities = kpi.get(
        "activity",
        {}
    )

    for activity, values in activities.items():

        if values.get(
            "dropout_rate",
            0
        ) >= 15:

            actions.append(
                f"Review {activity}, where the recorded "
                f"dropout rate is {values.get('dropout_rate', 0):.2f}%."
            )

    locations = (
        kpi
        .get("location", {})
        .get("locations", {})
    )

    for location, values in locations.items():

        if values.get(
            "dropout_rate",
            0
        ) >= 20:

            actions.append(
                f"Prioritise {location} for review because "
                f"its dropout rate is "
                f"{values.get('dropout_rate', 0):.2f}%."
            )

    if not actions:

        actions.append(
            "Continue routine monitoring of programme "
            "continuity, attendance, learning outcomes "
            "and data completeness."
        )

    for number, action in enumerate(
        actions,
        start=1
    ):

        add_numbered(
            document,
            number,
            action
        )


# ============================================================
# DATA QUALITY
# ============================================================

def add_data_quality(
    document,
    kpi
):

    add_heading(
        document,
        "13. Data Quality and Assessment Coverage",
        1
    )

    overall = kpi.get(
        "overall",
        {}
    )

    total = overall.get(
        "total_records",
        0
    )

    unique = overall.get(
        "unique_valid_software_ids",
        0
    )

    difference = (
        total - unique
    )

    add_paragraph(
        document,
        f"The dataset contains {total:,} records and "
        f"{unique:,} unique valid Software IDs. The "
        f"difference of {difference:,} records should be "
        f"reviewed as part of routine data-quality checks."
    )

    language = kpi.get(
        "language",
        {}
    )

    if language:

        add_paragraph(
            document,
            f"Language assessment coverage: "
            f"{language.get('total_assessed', 0):,} assessed; "
            f"{language.get('missing', 0):,} without a recorded "
            f"assessment."
        )

    mathematics = kpi.get(
        "mathematics",
        {}
    )

    if mathematics:

        add_paragraph(
            document,
            f"Mathematics assessment coverage: "
            f"{mathematics.get('total_assessed', 0):,} assessed; "
            f"{mathematics.get('missing', 0):,} without a "
            f"recorded assessment."
        )


# ============================================================
# CONCLUSION
# ============================================================

def add_conclusion(
    document,
    funder,
    kpi
):

    add_heading(
        document,
        "14. Conclusion",
        1
    )

    status = kpi.get(
        "status",
        {}
    )

    continuation = status.get(
        "continuation_rate",
        0
    )

    add_paragraph(
        document,
        f"The reporting period reflects the contribution "
        f"of {funder} towards the Community Education "
        f"Programme. The programme reached children through "
        f"multiple community-based activities and recorded "
        f"a continuation rate of {continuation:.2f}% in the "
        f"latest status period."
    )

    add_paragraph(
        document,
        "The findings highlight both areas of strong "
        "programme engagement and areas where continued "
        "attention can strengthen outcomes. Regular "
        "monitoring of attendance, continuity, migration "
        "and learning assessment coverage will help the "
        "programme respond to emerging needs and improve "
        "the quality of support provided to children."
    )


# ============================================================
# GENERATE ONE REPORT
# ============================================================

def generate_report(
    funder,
    kpi,
    insights
):

    document = Document()

    # Default font
    styles = document.styles

    styles["Normal"].font.name = (
        "Aptos"
    )

    styles["Normal"].font.size = Pt(10)

    create_cover(
        document,
        funder,
        kpi
    )

    add_programme_introduction(
        document
    )

    add_executive_summary(
        document,
        funder,
        kpi
    )

    add_key_highlights(
        document,
        kpi,
        insights
    )

    add_gender_section(
        document,
        kpi
    )

    add_status_section(
        document,
        kpi,
        insights
    )

    add_activity_section(
        document,
        kpi,
        insights
    )

    add_attendance_section(
        document,
        kpi,
        insights
    )

    add_language_section(
        document,
        kpi,
        insights
    )

    add_mathematics_section(
        document,
        kpi,
        insights
    )

    add_location_section(
        document,
        kpi,
        insights
    )

    add_interpretation(
        document,
        funder,
        kpi,
        insights
    )

    add_recommendations(
        document,
        funder,
        kpi,
        insights
    )

    add_data_quality(
        document,
        kpi
    )

    add_conclusion(
        document,
        funder,
        kpi
    )

    filename = (
        funder
        .replace("/", "_")
        .replace("\\", "_")
        .replace(" ", "_")
        + "_Progress_Report.docx"
    )

    output_file = os.path.join(
        OUTPUT_DIR,
        filename
    )

    document.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("========================================")
    print("FUNDER REPORT GENERATOR V3")
    print("========================================")

    if not os.path.exists(
        KPI_DIR
    ):

        print(
            "ERROR: Funder KPI folder not found."
        )

        return

    files = sorted(

        file

        for file in os.listdir(
            KPI_DIR
        )

        if file.endswith(
            "_kpis.json"
        )
        and file != "funder_summary.json"

    )

    print()
    print(
        f"Funder KPI files found: {len(files)}"
    )

    generated = []

    for filename in files:

        funder = filename.replace(
            "_kpis.json",
            ""
        )

        kpi_file = os.path.join(
            KPI_DIR,
            filename
        )

        insight_file = os.path.join(
            INSIGHT_DIR,
            f"{funder}_insights.json"
        )

        print()
        print(
            f"Generating report: {funder}"
        )

        with open(
            kpi_file,
            "r",
            encoding="utf-8"
        ) as file:

            kpi = json.load(
                file
            )

        if os.path.exists(
            insight_file
        ):

            with open(
                insight_file,
                "r",
                encoding="utf-8"
            ) as file:

                insights = json.load(
                    file
                )

        else:

            insights = {}

        output_file = generate_report(
            funder,
            kpi,
            insights
        )

        generated.append(
            output_file
        )

        print(
            f"Saved: {output_file}"
        )

    print()
    print("========================================")
    print("REPORT GENERATION COMPLETED")
    print("========================================")

    print()
    print(
        f"Reports generated: {len(generated)}"
    )

    print()
    print("Output folder:")
    print(OUTPUT_DIR)

    print()

    for file in generated:

        print(file)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()